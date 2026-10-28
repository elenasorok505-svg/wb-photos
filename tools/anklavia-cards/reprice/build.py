import json, sys, os, shutil, collections, warnings, statistics as st
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from load import load
import engine as E

D = '/home/claude/reprice/'
OUT = D + 'out/'
os.makedirs(OUT, exist_ok=True)

IR = json.load(open(D + 'ir.json'))          # поиск по фото: ключ "ds|группа" → {res, url}
WQ = json.load(open(D + 'wq.json'))          # поиск по словам: запрос → список цен

rows = load()

# ---------- цена аналога на WB для каждой строки ----------
def tmed(q):
    v = WQ.get(q) or []
    return st.median(v) if len(v) >= 10 else None

CAP = 1.5   # фото-поиск дороже медианы категории больше чем в 1,5 раза → нашёл другие товары (люстры, комплекты)

for r in rows:
    key = f"{r['ds']}|{r['group']}"
    ir = IR.get(key) or {}
    r['photo'] = E.analog_median(ir['res']) if ir.get('res') else None
    r['wb_url'] = ir.get('url', '')
    r['cat'] = (tmed(r['q']) if r['q'] else None) or (float(r['oldmed']) if r.get('oldmed') else None)

# постельное: фото-поиск показывает цену младшего размера; старшие размеры — по соотношению медиан размеров
by_g = collections.defaultdict(list)
for r in rows:
    if r['ds'] == 'bed': by_g[r['group']].append(r)
for g, rs in by_g.items():
    base = min((r['cat'] for r in rs if r['cat']), default=None)
    for r in rs:
        r['size_k'] = (r['cat'] / base) if (base and r['cat']) else 1.0

for r in rows:
    m = r['photo'] * r.get('size_k', 1.0) if r['photo'] else None
    if m is None:
        r['wb'], r['wb_src'] = r['cat'], 'медиана по категории (фото не нашлось)'
    elif r['cat'] and m > CAP * r['cat']:
        r['wb'], r['wb_src'] = r['cat'], 'медиана по категории (по фото — другие товары)'
    else:
        r['wb'] = m
        r['wb_src'] = 'по фото × размер' if r.get('size_k', 1.0) != 1.0 else 'поиск по фото'

bad = [r['vart'] for r in rows if not r['wb']]
if bad: print('нет цены WB:', bad[:20], len(bad))

for r in rows:
    r.update(E.price(r['cost'], r['w'], r['wb'], use_ceiling=r['wb_src'].startswith('медиана')))

json.dump([{k: v for k, v in r.items()} for r in rows], open(D + 'priced.json', 'w'), ensure_ascii=False, default=str)

# ---------- карточки: цена в юанях ----------
NEWCODE = {'bags200': ('ANK-BAG-', 'ANK-SUM-')}
NAMES = {'bags200': ('Сумки_200_карточки.xlsx', 'Сумки_200_расчёт_цен.xlsx', 'Сумки (200)'),
         'bags92': ('Сумки_92_карточки.xlsx', 'Сумки_92_расчёт_цен.xlsx', 'Сумки (92, старый файл)'),
         'lamps': ('Абажуры_200_карточки.xlsx', 'Абажуры_расчёт_цен.xlsx', 'Абажуры и плафоны'),
         'mas': ('Массажеры_100_карточки.xlsx', 'Массажеры_100_расчёт_цен.xlsx', 'Массажёры'),
         'bed': ('Постельное_100_карточки.xlsx', 'Постельное_100_расчёт_цен.xlsx', 'Постельное бельё')}

for ds, (card, calc, title) in NAMES.items():
    rs = [r for r in rows if r['ds'] == ds]
    wb = openpyxl.load_workbook(D + card)
    ws = wb.worksheets[0]
    H = [c.value for c in ws[1]]
    vcol = H.index('Артикул варианта') + 1
    for r in rs:
        assert str(ws.cell(r['crow'], vcol).value).strip() == r['vart'], (ds, r['vart'])
        ws.cell(r['crow'], 8).value = r['cny']
    if ds in NEWCODE:
        a, b = NEWCODE[ds]
        for i in range(3, ws.max_row + 1):
            for c in (1, 2, vcol):
                v = ws.cell(i, c).value
                if isinstance(v, str) and v.startswith(a): ws.cell(i, c).value = b + v[len(a):]
        for r in rs:
            r['vart_new'] = b + r['vart'][len(a):]; r['group_new'] = b + r['group'][len(a):]
    wb.save(OUT + card)

    # ---------- таблица расчёта ----------
    cw = openpyxl.Workbook()
    ps = cw.active; ps.title = 'Параметры'
    params = [('Курс ЦБ, ₽ за 1 ¥', E.P['rate'], 'ЦБ РФ на 28.10.2026 — обновляйте'),
              ('Комиссия маркетплейса', E.P['commission'], ''), ('Вывод денег', E.P['withdrawal'], ''),
              ('Минимальная прибыль с продажи, ₽', E.P['min_profit'], 'ниже этой прибыли цену не опускаем'),
              ('Насколько дешевле WB', E.P['disc'], 'цена = цена этого товара на WB − 15%'),
              ('Тариф 1: ¥ за кг', E.P['t1k'], 'логистика: берём дешевле из двух тарифов'),
              ('Тариф 1: ¥ за заказ', E.P['t1o'], ''), ('Тариф 2: ¥ за кг', E.P['t2k'], ''),
              ('Тариф 2: ¥ за заказ', E.P['t2o'], ''), ('Упаковка до 300 г, г', E.P['pack_l'], ''),
              ('Упаковка от 300 г, г', E.P['pack_h'], ''),
              ('Множитель: коэффициент a', 8.34, 'потолок нужен только когда точного аналога на WB нет'),
              ('Множитель: степень b', 0.222, 'K = a × закупка^(−b)'), ('Множитель: минимум', 1.6, ''), ('Множитель: максимум', 6, '')]
    BLUE = Font(name='Arial', color='0000FF'); B = Font(name='Arial', bold=True); A = Font(name='Arial')
    for i, (k, v, c) in enumerate(params, start=1):
        ps.cell(i, 1, k).font = A; x = ps.cell(i, 2, v); x.font = BLUE; ps.cell(i, 3, c).font = A
        if isinstance(v, float) and v < 1: x.number_format = '0%'
    notes = ['', 'Как считается цена',
             '1. Цена этого товара на WB — медиана цен первых 8 совпадений в поиске WB по фото товара (ссылка на поиск — в таблице).',
             '2. Рынок = цена на WB − 15%.',
             '3. Себестоимость = (закупка + логистика) × курс. Пол = (себестоимость + 150 ₽) / (0,9 × 0,99) — после комиссий остаётся ≥150 ₽ прибыли.',
             '4. Если точного аналога по фото нет (или по фото нашлись другие товары — люстры, наборы), берётся медиана категории, но не выше потолка: (закупка × K + логистика) / 0,891 × курс.',
             '5. Цена в юанях = большее из: рынок / курс (вниз до целого ¥) и пол / курс (вверх до целого ¥).',
             '6. В шаблон Анклавии идёт цена в юанях (колонка «Цена»): кабинет хранит цены в ¥.']
    for i, t in enumerate(notes, start=len(params) + 1):
        ps.cell(i, 1, t).font = B if 'Как' in t else A
    ps.column_dimensions['A'].width = 36; ps.column_dimensions['C'].width = 60

    s = cw.create_sheet('Расчёт', 0)
    hdr = ['Группа', 'Артикул варианта', 'Название', 'Вариант', 'Закупка, ¥', 'Вес товара, г',
           'Вес для логистики, кг', 'Логистика, ¥', 'Себестоимость, ₽', 'Пол (прибыль 150 ₽), ₽',
           'Цена этого товара на WB, ₽', 'Откуда цена WB', 'Поиск на WB', 'Рынок (WB −15%), ₽',
           'ЦЕНА, ¥', 'Цена в рублях, ₽', 'Вы получаете, ₽', 'Прибыль, ₽', 'Наценка, x', 'Дешевле WB на',
           'Основа цены', 'Закупка: источник', 'В файле карточек, ¥', 'Совпадает', 'WB по фото (медиана 8), ₽', 'Медиана категории WB, ₽']
    for j, h in enumerate(hdr, 1):
        c = s.cell(1, j, h); c.font = B; c.alignment = Alignment(wrap_text=True, vertical='top')
        c.fill = PatternFill('solid', fgColor='EFE7DC')
    s.row_dimensions[1].height = 45
    pr = "Параметры!$B$"
    for i, r in enumerate(rs, start=2):
        vals = [r.get('group_new', r['group']), r.get('vart_new', r['vart']), r['name'], r['var'], r['cost'], r['w']]
        for j, v in enumerate(vals, 1): s.cell(i, j, v).font = BLUE if j in (5, 6) else A
        s.cell(i, 7, f'=ROUNDUP((F{i}+IF(F{i}<300,{pr}10,{pr}11))/100,0)/10')
        s.cell(i, 8, f'=MIN({pr}6*G{i}+{pr}7,{pr}8*G{i}+{pr}9)')
        s.cell(i, 9, f'=(E{i}+H{i})*{pr}1')
        s.cell(i, 10, f'=(I{i}+{pr}4)/((1-{pr}2)*(1-{pr}3))')
        s.cell(i, 11, round(r['wb'], 2)).font = BLUE
        s.cell(i, 12, r['wb_src'])
        if r['wb_url']:
            c = s.cell(i, 13, 'открыть'); c.hyperlink = r['wb_url']; c.font = Font(name='Arial', color='0563C1', underline='single')
        ceil = f'(E{i}*MAX({pr}14,MIN({pr}15,{pr}12*E{i}^(-{pr}13)))+H{i})/((1-{pr}2)*(1-{pr}3))*{pr}1'
        s.cell(i, 14, f'=IF(LEFT(L{i},7)="медиана",MIN(K{i}*(1-{pr}5),{ceil}),K{i}*(1-{pr}5))')
        s.cell(i, 15, f'=MAX(ROUNDUP(J{i}/{pr}1,0),INT(N{i}/{pr}1))').font = B
        s.cell(i, 16, f'=O{i}*{pr}1')
        s.cell(i, 17, f'=P{i}*(1-{pr}2)*(1-{pr}3)')
        s.cell(i, 18, f'=Q{i}-I{i}')
        s.cell(i, 19, f'=P{i}/I{i}')
        s.cell(i, 20, f'=1-P{i}/K{i}')
        s.cell(i, 21, f'=IF(INT(N{i}/{pr}1)<ROUNDUP(J{i}/{pr}1,0),"мин. прибыль",IF(N{i}<K{i}*(1-{pr}5)-0.01,"потолок (нет точного аналога)","WB −15%"))')
        s.cell(i, 22, 'закупка 1688' if r['cost_src'] == '1688' else 'оценка по прошлому расчёту')
        s.cell(i, 23, r['cny']).font = BLUE
        s.cell(i, 24, f'=IF(O{i}=W{i},"да","НЕТ")')
        if r['photo']: s.cell(i, 25, round(r['photo'] * r.get('size_k', 1.0))).number_format = '#,##0'
        if r['cat']: s.cell(i, 26, round(r['cat'])).number_format = '#,##0'
        for j in (9, 10, 11, 14, 16, 17, 18): s.cell(i, j).number_format = '#,##0'
        s.cell(i, 8).number_format = '0.0'; s.cell(i, 19).number_format = '0.00"x"'
        s.cell(i, 20).number_format = '0%'
    widths = [13, 17, 40, 16, 9, 9, 10, 9, 11, 11, 12, 22, 9, 11, 9, 10, 11, 10, 9, 9, 13, 16, 10, 9, 12, 12]
    for j, w in enumerate(widths, 1): s.column_dimensions[L(j)].width = w
    s.freeze_panes = 'C2'
    cw.save(OUT + calc)
    n = len(rs)
    print(f"{title}: {n} строк, медиана цены {st.median(r['cny'] for r in rs)} ¥ "
          f"(было {st.median(r['old'] for r in rs)} ₽ в колонке ¥), по полу: {sum(r['basis']!='WB −15%' for r in rs)}")
