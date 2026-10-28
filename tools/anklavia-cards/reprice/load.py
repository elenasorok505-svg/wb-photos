"""Собирает строки всех таблиц в один формат для пересчёта цен."""
import openpyxl, math, re, warnings
warnings.filterwarnings('ignore')
D = '/home/claude/reprice/'

OLD_RATE = 12.5764


def _old_ceiling(c, w, rate=OLD_RATE):
    pack = 50 if w < 300 else 80
    kg = math.ceil((w + pack) / 100) / 10
    log = min(58 * kg + 2, 43 * kg + 8)
    k = max(1.6, min(6.0, 8.34 * c ** (-0.222)))
    return (c * k + log) / 0.891 * rate


def invert_cost(old_price, w):
    """Закупка, при которой старый «потолок» давал эту цену (верх окна округления)."""
    step = 10 if old_price < 1000 else 100
    target = old_price + step
    lo, hi = 0.5, 2000
    for _ in range(60):
        mid = (lo + hi) / 2
        if _old_ceiling(mid, w) < target: lo = mid
        else: hi = mid
    return round(hi, 1)


def card_rows(fname, vcol):
    ws = openpyxl.load_workbook(D + fname, read_only=True).worksheets[0]
    out = {}
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < 3 or not r[0]: continue
        out[str(r[vcol]).strip()] = dict(row=i, group=r[0], price=r[7], name=r[2])
    return out


Q92 = {'ANK-BAG-001': 'сумка шоппер женская', 'ANK-BAG-002': 'сумка багет женская', 'ANK-BAG-003': 'сумка вечерняя со стразами',
       'ANK-BAG-004': 'сумка полумесяц женская', 'ANK-BAG-005': 'сумка женская маленькая', 'ANK-BAG-006': 'сумка женская на цепочке',
       'ANK-BAG-007': 'сумка тоут женская', 'ANK-BAG-008': 'сумка хобо женская', 'ANK-BAG-009': 'сумка тоут женская',
       'ANK-BAG-010': 'сумка мешок женская', 'ANK-BAG-011': 'шоппер тканевый', 'ANK-BAG-012': 'сумка замшевая женская',
       'ANK-BAG-013': 'сумка бостон женская', 'ANK-BAG-014': 'сумка женская через плечо', 'ANK-BAG-015': 'сумка багет женская',
       'ANK-BAG-016': 'сумка стеганая женская', 'ANK-BAG-017': 'сумка стеганая женская', 'ANK-BAG-018': 'сумка замшевая женская',
       'ANK-BAG-019': 'поясная сумка женская', 'ANK-BAG-020': 'сумка для телефона женская'}


def load():
    rows = []
    # сумки 200 и абажуры: одинаковые таблицы расчёта
    for ds, calc, card in [('bags200', 'Сумки_расчёт_цен.xlsx', 'Сумки_200_карточки.xlsx'),
                           ('lamps', 'Абажуры_расчёт_цен.xlsx', 'Абажуры_200_карточки.xlsx')]:
        cr = card_rows(card, 10)
        ws = openpyxl.load_workbook(D + calc, read_only=True).worksheets[0]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r[1] or r[1] not in cr: continue
            rows.append(dict(ds=ds, card=card, crow=cr[r[1]]['row'], group=r[0], vart=r[1], name=r[2], var=r[3],
                             cost=float(r[4]), w=float(r[5]), q=r[21], old=cr[r[1]]['price'], cost_src='1688', oldmed=r[12]))
    # массажёры и постельное
    for ds, calc, card, vcol in [('mas', 'Массажеры_100_расчёт_цен.xlsx', 'Массажеры_100_карточки.xlsx', 10),
                                 ('bed', 'Постельное_100_расчёт_цен.xlsx', 'Постельное_100_карточки.xlsx', 11)]:
        cr = card_rows(card, vcol)
        ws = openpyxl.load_workbook(D + calc, read_only=True).worksheets[0]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r[1] or r[1] == '—' or not isinstance(r[5], (int, float)): continue
            v = re.sub(r'\s*/\s*вариант\s*', '-', str(r[1]).strip())
            if v not in cr: continue
            rows.append(dict(ds=ds, card=card, crow=cr[v]['row'], group=cr[v]['group'], vart=v, name=r[2], var=r[3],
                             cost=float(r[5]), w=float(r[6]), q=r[7], old=cr[v]['price'], cost_src='1688', oldmed=r[8]))
        miss = set(cr) - {x['vart'] for x in rows if x['ds'] == ds}
        if miss: print('карточки без расчёта', ds, sorted(miss)[:10], len(miss))
    # сумки 92: таблицы расчёта нет
    ws = openpyxl.load_workbook(D + 'Сумки_92_карточки.xlsx', read_only=True).worksheets[0]
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < 3 or not r[0]: continue
        wp = float(r[11]); w = wp - 50 if wp - 50 < 300 else wp - 80
        g = r[1][:11]
        if g == 'ANK-BAG-002':
            cost, src = 14.5, '1688'
        else:
            cost, src = invert_cost(r[7], w), 'оценка'
        rows.append(dict(ds='bags92', card='Сумки_92_карточки.xlsx', crow=i, group=g, vart=r[1], name=r[2], var=r[9],
                         cost=cost, w=w, q=Q92.get(g), old=r[7], cost_src=src, oldmed=None))
    return rows


if __name__ == '__main__':
    import collections
    R = load()
    print(collections.Counter(r['ds'] for r in R))
    for ds in ['bags200', 'lamps', 'mas', 'bed', 'bags92']:
        print([ (r['vart'], r['cost'], r['w'], r['q'], r['old']) for r in R if r['ds'] == ds][:2])
