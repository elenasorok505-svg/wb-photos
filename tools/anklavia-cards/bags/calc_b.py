"""Таблица расчёта цен (формулы Excel повторяют pricing_logic.calc_price). python3 calc_l.py out.xlsx"""
import sys, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, '/home/claude/bed')
from pricing_logic import PARAMS

def main(out):
    rows = json.load(open('/home/claude/bags/rows.json'))
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Расчёт цен'
    par = wb.create_sheet('Параметры')
    pp = [('Курс ¥→₽', PARAMS['rate_cny_rub']), ('Комиссия маркетплейса', PARAMS['commission']), ('Комиссия за вывод', PARAMS['withdrawal']),
          ('Мин. прибыль, ₽', PARAMS['min_profit_rub']), ('Скидка к медиане WB', PARAMS['wb_discount']), ('k_a', PARAMS['k_a']), ('k_b', PARAMS['k_b']),
          ('K мин', PARAMS['k_min']), ('K макс', PARAMS['k_max']), ('Тариф 1, ¥/кг', 58), ('Тариф 1, ¥/заказ', 2), ('Тариф 2, ¥/кг', 43), ('Тариф 2, ¥/заказ', 8),
          ('Упаковка до 300 г, г', 50), ('Упаковка от 300 г, г', 80)]
    for i, (a, b) in enumerate(pp, 1): par.cell(i, 1, a); par.cell(i, 2, b)
    par.column_dimensions['A'].width = 28
    R = {k: f"Параметры!$B${i}" for i, k in enumerate(['rate', 'com', 'wd', 'minp', 'disc', 'ka', 'kb', 'kmin', 'kmax', 't1', 'o1', 't2', 'o2', 'pl', 'ph'], 1)}
    keep = f"((1-{R['com']})*(1-{R['wd']}))"
    hdr = ['Группа', 'Артикул варианта', 'Название', 'Цвет', 'Закупка, ¥', 'Вес товара, г', 'Вес для логистики, кг', 'Логистика, ¥', 'Себестоимость, ₽',
           'Множитель K', 'Потолок, ₽', 'Пол (мин. прибыль), ₽', 'Медиана WB, ₽', 'Рынок (WB −15%), ₽', 'Цена, ₽', 'Цена по формуле, ₽', 'Совпадает',
           'Вы получаете, ₽', 'Прибыль, ₽', 'Наценка, x', 'Основа цены', 'Запрос WB', 'Размеры']
    for j, h in enumerate(hdr, 1):
        c = ws.cell(1, j, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='5B4636'); c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, x in enumerate(rows, 2):
        ws.cell(i, 1, x['group']); ws.cell(i, 2, x['art']); ws.cell(i, 3, x['title']); ws.cell(i, 4, x['color'])
        ws.cell(i, 5, x['cost']); ws.cell(i, 6, x['g'])
        ws.cell(i, 7, f"=ROUNDUP((F{i}+IF(F{i}<300,{R['pl']},{R['ph']}))/100,0)/10")
        ws.cell(i, 8, f"=MIN({R['t1']}*G{i}+{R['o1']},{R['t2']}*G{i}+{R['o2']})")
        ws.cell(i, 9, f"=(E{i}+H{i})*{R['rate']}")
        ws.cell(i, 10, f"=MAX({R['kmin']},MIN({R['kmax']},{R['ka']}*E{i}^(-{R['kb']})))")
        ws.cell(i, 11, f"=(E{i}*J{i}+H{i})/{keep}*{R['rate']}")
        ws.cell(i, 12, f"=I{i}/{keep}+{R['minp']}")
        if x['med']: ws.cell(i, 13, x['med']); ws.cell(i, 14, f"=M{i}*(1-{R['disc']})")
        ws.cell(i, 15, x['price'])
        raw = f'IF(N{i}="",K{i},MIN(K{i},N{i}))'
        pr = f'IF({raw}<1000,INT((INT({raw})+1)/10)*10-1,INT((INT({raw})+10)/100)*100-10)'
        fl0 = f'IF(L{i}<1000,ROUNDUP(L{i}/10,0)*10-1,ROUNDUP(L{i}/100,0)*100-10)'
        fl = f'IF({fl0}<L{i},{fl0}+IF({fl0}<1000,10,100),{fl0})'
        ws.cell(i, 16, f'=IF({pr}<L{i},{fl},{pr})')
        ws.cell(i, 17, f'=IF(O{i}=P{i},"да","НЕТ")')
        ws.cell(i, 18, f"=O{i}*{keep}"); ws.cell(i, 19, f"=R{i}-I{i}"); ws.cell(i, 20, f"=O{i}/I{i}")
        ws.cell(i, 21, x['basis']); ws.cell(i, 22, x['q']); ws.cell(i, 23, 'примерные — уточнить' if x.get('approx') else 'от поставщика')
        for j in (9, 11, 12, 14, 18, 19): ws.cell(i, j).number_format = '0'
        ws.cell(i, 10).number_format = '0.00'; ws.cell(i, 20).number_format = '0.00'
    widths = [13, 17, 50, 18, 10, 10, 11, 11, 13, 10, 11, 13, 12, 13, 10, 12, 10, 12, 11, 10, 18, 30, 20]
    for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = f'A1:{L(len(hdr))}{len(rows) + 1}'
    wb.save(out); print('rows', len(rows))

if __name__ == '__main__':
    main(sys.argv[1])
