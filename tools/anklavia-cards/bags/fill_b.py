"""Шаблон Anklavia «Сумки женские»: товар = Группа (Артикул продавца = Группа), цвета = строки-варианты. python3 fill_b.py out.xlsx"""
import sys, json, os
import openpyxl
BASE = 'https://raw.githubusercontent.com/elenasorok505-svg/wb-photos/main/anklavia/bags'
OUT = '/home/claude/bags/out'
TNVED = {'pu': '4202221000', 'pu_matte': '4202221000', 'pu_lacquer': '4202221000', 'straw': '4202290000', 'crystal': '4202290000'}
CARE = {'pu': 'Уход: протирайте влажной салфеткой, не сушите на батарее.', 'pu_matte': 'Уход: протирайте влажной салфеткой, не сушите на батарее.',
        'pu_lacquer': 'Уход: протирайте мягкой влажной тканью, храните отдельно от светлых вещей.',
        'suede': 'Уход: чистите сухой мягкой щёткой по направлению ворса, берегите от сильной влаги.',
        'fur': 'Уход: расчёсывайте мех мягкой щёткой, не стирайте в машине, сушите при комнатной температуре.',
        'straw': 'Уход: удаляйте пыль мягкой кистью, берегите от сильной влаги и не сминайте.',
        'crystal': 'Уход: протирайте сухой мягкой тканью, храните в мешочке, берегите от ударов.',
        'sequin': 'Уход: протирайте сухой мягкой тканью, храните в мешочке.',
        'satin': 'Уход: чистите сухой мягкой щёткой, берегите от влаги и пятен.', 'velvet': 'Уход: чистите сухой мягкой щёткой по ворсу.'}
CARE_TEX = 'Уход: удаляйте пятна влажной салфеткой с мягким средством, сушите естественным образом.'
FEAT = {'cross': 'Регулируемый ремень', 'shoulder': 'Удобные ручки', 'back': 'Трансформер', 'clutch': 'Компактная', 'waist': 'Регулируемый ремень', 'hand': 'Удобные ручки'}

def fmt(x): return str(int(x)) if float(x).is_integer() else str(x).replace('.', ',')
def low(s): return s[0].lower() + s[1:] if s else s

def kit(p):
    k = ['Рюкзак×1' if p['type'] == 'backpack' else 'Сумка×1']
    if p['strap']: k.append('Плечевой ремень×1')
    if p['decor'] == 'Брелок': k.append('Брелок×1')
    if p['decor'] == 'Картхолдер': k.append('Картхолдер×1')
    return '; '.join(k)

def descr(p, v):
    d = p['dims']; ps = [f"{p['title']}. Цвет: {v['color']}."]
    ps.append(f"{p['scene_t']}. {p['fit_s']}: {', '.join(low(x) for x in p['fits'][:6])}.")
    ps.append(' '.join(f"{a}: {low(b)}." for _, a, b in p['feats'][:4]))
    ps.append(f"Материал: {p['mat']}. Подкладка: {low(p['lining'])}. Размер: {fmt(d['w'])}×{fmt(d['h'])}×{fmt(d['d'])} см (ширина×высота×глубина).")
    ps.append(f"Застёжка: {low(p['closure'])}. Отделений: {p['comp']}.")
    ps.append(p['strap_txt'] + '.')
    ps.append(CARE.get(p['line'], CARE_TEX))
    ps.append(f"В комплекте: {kit(p).replace('×1', '').replace(';', ',').lower()}.")
    return ' '.join(ps)[:3900]

def main(out_xlsx):
    P = json.load(open('/home/claude/bags/prods.json'))
    wb = openpyxl.load_workbook('/home/claude/bags/tpl.xlsx'); ws = wb['Товары']
    hdr = [c.value for c in ws[1]]; col = {h: i + 1 for i, h in enumerate(hdr)}
    r = 3; rows = []; titles = set()
    for p in P:
        t = p['title'][:120]
        if t in titles: raise SystemExit('dup title ' + t)
        titles.add(t); d = p['dims']; eve = p['kind'] == 'Вечерние'
        feats = [FEAT[p['wear'][0]]] + (['Съёмный ремень'] if p['strap'] else []) + (['Вместительная'] if p['size'] in ('large', 'xl') else ['Лёгкая'])
        gpack = int(p['g'] + 60)
        for v in p['variants']:
            k = v['key']; dd = f'{OUT}/{k}'
            files = sorted(f for f in os.listdir(dd) if f.endswith('.jpg'))[:10] if os.path.isdir(dd) else []
            pk = v['pkg']
            val = {'Группа': p['art'], 'Артикул продавца': p['art'], 'Название': t, 'Описание': descr(p, v),
                   'Фото': ';'.join(f'{BASE}/{v["vart"]}/{f}' for f in files), 'Цена, ₽': v['price'], 'Цвет': v['color'], 'Артикул варианта': v['vart'],
                   'Вес с упаковкой, г': gpack, 'Длина упаковки, см': pk[0], 'Ширина упаковки, см': pk[1], 'Высота упаковки, см': pk[2],
                   'Страна производства': 'CN', 'Код ТН ВЭД': TNVED.get(p['line'], '4202229000'), 'Упаковка': 'Пакет', 'Комплектация': kit(p),
                   'Возрастное ограничение': 0, 'Вид': p['kind'], 'Вид застёжки': p['closure'], 'Декор': p['decor'], 'Пол': 'Женский',
                   'Рисунок': p['pattern'], 'Карманы': p['pockets'], 'Материал подкладки': p['lining'], 'Особенности': ';'.join(feats[:3]),
                   'Назначение ремня': ';'.join(p['strap_purpose'][:3]), 'Фактура': p['texture'], 'Фасон': p['cut_style'],
                   'Количество отделений': p['comp'], 'Материал': p['mat_attr'], 'Длина плечевого ремня': p['strap'] or None,
                   'Глубина': d['d'], 'Высота': d['h'], 'Ширина': d['w'],
                   'Повод': 'Свадьба;Выпускной;День рождения' if eve else 'День рождения;8 Марта;Без повода', 'Кому в подарок': 'Женщине;Подруге;Маме'}
            for kk, x in val.items():
                if kk not in col or x is None: continue
                cell = ws.cell(r, col[kk], x)
                if kk == 'Код ТН ВЭД': cell.number_format = '@'
            c = v['calc']
            rows.append(dict(group=p['art'], art=v['vart'], title=t, color=v['color'], line=p['line'], key=k, cost=v['cost'], g=v['g'], med=v.get('med'),
                             price=v['price'], floor=c.get('floor_rub'), market=c.get('market_rub'), profit=c.get('profit_rub'), basis=c.get('basis'),
                             photos=len(files), q=p['q'], approx=p.get('approx', False)))
            r += 1
    wb.save(out_xlsx)
    json.dump(rows, open('/home/claude/bags/rows.json', 'w'), ensure_ascii=False, indent=1)
    print('rows', len(rows), 'products', len(P), 'no photos', sum(1 for x in rows if x['photos'] == 0))

if __name__ == '__main__':
    main(sys.argv[1])
